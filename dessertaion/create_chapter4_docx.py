import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="8" w:space="0" w:color="333333"/>
                <w:bottom w:val="single" w:sz="8" w:space="0" w:color="333333"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def build_docx():
    doc = docx.Document()

    # Page setup - Margins 1 inch
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run("Chapter 4 – Results")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy accent

    # Read verified corrected markdown text
    with open("dessertaion/chapter4_results_corrected.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    # Parse sections from markdown
    sec_blocks = md_text.split("## ")

    for block in sec_blocks[1:]:
        lines = block.strip().split("\n")
        header = lines[0].strip()
        body_paras = [p.strip() for p in "\n".join(lines[1:]).split("\n\n") if p.strip()]

        # Heading 2
        h2 = doc.add_heading(level=2)
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(6)
        h2_run = h2.add_run(header)
        h2_run.font.name = 'Calibri'
        h2_run.font.size = Pt(15)
        h2_run.font.bold = True
        h2_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        # Add text paragraphs and embed corresponding tables/figures
        if header.startswith("4.1"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)
            
            # Table 4.1: Experimental Coverage
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(4)
            r_cap = p_cap.add_run("Table 4.1. Experimental Coverage Summary (Run3)")
            r_cap.bold = True
            r_cap.font.size = Pt(10)

            t1 = doc.add_table(rows=4, cols=8)
            t1.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Dataset", "Model", "Expected Runs", "Actual Runs", "Completed", "Failed", "Coverage Status", "Attributable"]
            for i, h in enumerate(headers):
                cell = t1.cell(0, i)
                cell.text = h
                set_cell_background(cell, "1B365D")
                set_cell_margins(cell)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9.5)

            rows_data = [
                ["Rotten Tomatoes", "Gemma 2 2B", "26", "36", "34", "2", "Executed (34 Complete / 2 Failed)", "Yes (RT)"],
                ["Rotten Tomatoes", "Qwen 2.5 3B", "26", "42", "38", "4", "Executed (38 Complete / 4 Failed)", "Yes (RT)"],
                ["DBPedia 14", "DeepSeek R1 1.5B", "26", "22", "20", "2", "Executed (20 Complete / 2 Failed)", "Yes (DBPedia)"]
            ]
            for r_idx, rdata in enumerate(rows_data, start=1):
                bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx, val in enumerate(rdata):
                    cell = t1.cell(r_idx, c_idx)
                    cell.text = val
                    set_cell_background(cell, bg)
                    set_cell_margins(cell)
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_table_borders(t1)

        elif header.startswith("4.2"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)

            # Table 4.2: Overall Performance
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(4)
            r_cap = p_cap.add_run("Table 4.2. Overall Performance Summary Across Model–Dataset Conditions")
            r_cap.bold = True
            r_cap.font.size = Pt(10)

            t2 = doc.add_table(rows=4, cols=8)
            t2.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Dataset", "Model", "Completed Runs", "Mean Accuracy", "Std Acc", "Mean F1", "Mean Latency", "Token Status"]
            for i, h in enumerate(headers):
                cell = t2.cell(0, i)
                cell.text = h
                set_cell_background(cell, "1B365D")
                set_cell_margins(cell)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9.5)

            t2_data = [
                ["DBPedia 14", "DeepSeek R1 1.5B", "20", "0.3295", "0.1311", "0.1932", "3342.0 ms", "Unrecorded"],
                ["Rotten Tomatoes", "Gemma 2 2B", "34", "0.2265", "0.2486", "0.1200", "1377.0 ms", "444.5 total"],
                ["Rotten Tomatoes", "Qwen 2.5 3B", "38", "0.0421", "0.1154", "0.0159", "913.3 ms", "Unrecorded"]
            ]
            for r_idx, rdata in enumerate(t2_data, start=1):
                bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx, val in enumerate(rdata):
                    cell = t2.cell(r_idx, c_idx)
                    cell.text = val
                    set_cell_background(cell, bg)
                    set_cell_margins(cell)
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_table_borders(t2)

        elif header.startswith("4.3"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)

            # Insert Figures 4.1 and 4.2
            fig_p = doc.add_paragraph()
            fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p.paragraph_format.space_before = Pt(8)
            fig_p.paragraph_format.space_after = Pt(2)
            run = fig_p.add_run()
            run.add_picture("Run3/figures/01_prompt_structure_accuracy.png", width=Inches(5.5))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            r_cap = p_cap.add_run("Figure 4.1. Mean Accuracy by Prompt Structure (Instruction vs. Example-Based vs. Mixed at 3-shot)")
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True

            fig_p2 = doc.add_paragraph()
            fig_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p2.paragraph_format.space_before = Pt(4)
            fig_p2.paragraph_format.space_after = Pt(2)
            run2 = fig_p2.add_run()
            run2.add_picture("Run3/figures/02_prompt_structure_f1.png", width=Inches(5.5))
            
            p_cap2 = doc.add_paragraph()
            p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap2.paragraph_format.space_after = Pt(10)
            r_cap2 = p_cap2.add_run("Figure 4.2. Mean F1-Score by Prompt Structure Across Evaluated Model–Dataset Pairs")
            r_cap2.font.size = Pt(9.5)
            r_cap2.font.italic = True

        elif header.startswith("4.4"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)

            # Insert Figures 4.3 and 4.4
            fig_p = doc.add_paragraph()
            fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p.paragraph_format.space_before = Pt(8)
            fig_p.paragraph_format.space_after = Pt(2)
            run = fig_p.add_run()
            run.add_picture("Run3/figures/03_example_count_accuracy.png", width=Inches(5.5))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            r_cap = p_cap.add_run("Figure 4.3. Mean Accuracy Across Example Counts (0-shot, 1-shot, 3-shot, 5-shot Trajectories)")
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True

            fig_p2 = doc.add_paragraph()
            fig_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p2.paragraph_format.space_before = Pt(4)
            fig_p2.paragraph_format.space_after = Pt(2)
            run2 = fig_p2.add_run()
            run2.add_picture("Run3/figures/04_example_count_f1.png", width=Inches(5.5))
            
            p_cap2 = doc.add_paragraph()
            p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap2.paragraph_format.space_after = Pt(10)
            r_cap2 = p_cap2.add_run("Figure 4.4. Mean F1-Score Scaling Across Example Counts")
            r_cap2.font.size = Pt(9.5)
            r_cap2.font.italic = True

        elif header.startswith("4.5"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)

            # Table 4.3: Example Selection
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(4)
            r_cap = p_cap.add_run("Table 4.3. Example Selection Strategy Comparison (Random vs. Balanced vs. Sequential)")
            r_cap.bold = True
            r_cap.font.size = Pt(10)

            t3 = doc.add_table(rows=4, cols=6)
            t3.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Model", "Dataset", "Random Acc (F1)", "Balanced Acc (F1)", "Sequential Acc (F1)", "Failed Execution Runs"]
            for i, h in enumerate(headers):
                cell = t3.cell(0, i)
                cell.text = h
                set_cell_background(cell, "1B365D")
                set_cell_margins(cell)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9.5)

            t3_data = [
                ["DeepSeek R1 1.5B", "DBPedia 14", "0.3847 (0.2668)", "0.2111 (0.1162)", "Execution Failed", "2 (Vector Index Unindexed)"],
                ["Gemma 2 2B", "Rotten Tomatoes", "0.1167 (0.0571)", "0.1500 (0.0685)", "Unexecuted", "2 (Random Replications)"],
                ["Qwen 2.5 3B", "Rotten Tomatoes", "0.0167 (0.0076)", "0.0000 (0.0000)", "0.1000 (0.0455)", "2 (Sequential Retrieval)"]
            ]
            for r_idx, rdata in enumerate(t3_data, start=1):
                bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx, val in enumerate(rdata):
                    cell = t3.cell(r_idx, c_idx)
                    cell.text = val
                    set_cell_background(cell, bg)
                    set_cell_margins(cell)
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_table_borders(t3)

            # Insert Figure 4.5
            fig_p = doc.add_paragraph()
            fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p.paragraph_format.space_before = Pt(8)
            fig_p.paragraph_format.space_after = Pt(2)
            run = fig_p.add_run()
            run.add_picture("Run3/figures/07_example_selection_accuracy.png", width=Inches(5.5))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run("Figure 4.5. Mean Accuracy Comparison Across Example Selection Methods")
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True

        elif header.startswith("4.6"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)

            # Table 4.4: Example Ordering
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(4)
            r_cap = p_cap.add_run("Table 4.4. Example Ordering Strategy Comparison (Original vs. Random vs. Similarity)")
            r_cap.bold = True
            r_cap.font.size = Pt(10)

            t4 = doc.add_table(rows=4, cols=5)
            t4.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Model", "Dataset", "Original Acc (F1)", "Random Acc (F1)", "Similarity Acc (F1)"]
            for i, h in enumerate(headers):
                cell = t4.cell(0, i)
                cell.text = h
                set_cell_background(cell, "1B365D")
                set_cell_margins(cell)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9.5)

            t4_data = [
                ["DeepSeek R1 1.5B", "DBPedia 14", "0.3847 (0.2668)", "0.4097 (0.1935)", "0.3125 (0.1409)"],
                ["Gemma 2 2B", "Rotten Tomatoes", "0.1167 (0.0571)", "Unexecuted", "Unexecuted"],
                ["Qwen 2.5 3B", "Rotten Tomatoes", "0.0167 (0.0076)", "0.0500 (0.0227)", "0.2500 (0.0833)"]
            ]
            for r_idx, rdata in enumerate(t4_data, start=1):
                bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx, val in enumerate(rdata):
                    cell = t4.cell(r_idx, c_idx)
                    cell.text = val
                    set_cell_background(cell, bg)
                    set_cell_margins(cell)
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_table_borders(t4)

            # Insert Figure 4.6
            fig_p = doc.add_paragraph()
            fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p.paragraph_format.space_before = Pt(8)
            fig_p.paragraph_format.space_after = Pt(2)
            run = fig_p.add_run()
            run.add_picture("Run3/figures/09_example_ordering_accuracy.png", width=Inches(5.5))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run("Figure 4.6. Mean Accuracy Comparison Across Example Ordering Strategies")
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True

        elif header.startswith("4.7"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)

            # Table 4.5: Efficiency Results
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(4)
            r_cap = p_cap.add_run("Table 4.5. Computational Efficiency Summary (Latency & Token Metrics)")
            r_cap.bold = True
            r_cap.font.size = Pt(10)

            t5 = doc.add_table(rows=4, cols=6)
            t5.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Model", "Dataset", "Mean Latency", "Min Latency", "Max Latency", "Total Tokens (Input / Output)"]
            for i, h in enumerate(headers):
                cell = t5.cell(0, i)
                cell.text = h
                set_cell_background(cell, "1B365D")
                set_cell_margins(cell)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9.5)

            t5_data = [
                ["DeepSeek R1 1.5B", "DBPedia 14", "3342.0 ms", "3089.1 ms (5-shot)", "3561.3 ms (Similarity)", "Unrecorded in Run Table"],
                ["Gemma 2 2B", "Rotten Tomatoes", "1377.0 ms", "1164.2 ms (Mixed 3-shot)", "1932.3 ms (Balanced)", "444.5 tokens (161.9 / 282.6)"],
                ["Qwen 2.5 3B", "Rotten Tomatoes", "913.3 ms", "392.7 ms (5-shot)", "2604.2 ms (Sequential)", "Unrecorded in Run Table"]
            ]
            for r_idx, rdata in enumerate(t5_data, start=1):
                bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx, val in enumerate(rdata):
                    cell = t5.cell(r_idx, c_idx)
                    cell.text = val
                    set_cell_background(cell, bg)
                    set_cell_margins(cell)
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_table_borders(t5)

            # Insert Figure 4.7
            fig_p = doc.add_paragraph()
            fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p.paragraph_format.space_before = Pt(8)
            fig_p.paragraph_format.space_after = Pt(2)
            run = fig_p.add_run()
            run.add_picture("Run3/figures/14_latency_comparison.png", width=Inches(5.5))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run("Figure 4.7. Response Latency Comparison Across Evaluated Models and Tasks")
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True

        elif header.startswith("4.8"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)

            # Table 4.6: Reliability Summary
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(4)
            r_cap = p_cap.add_run("Table 4.6. Reliability & Replicability Variance Summary Across Configurations")
            r_cap.bold = True
            r_cap.font.size = Pt(10)

            t6 = doc.add_table(rows=5, cols=7)
            t6.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Model", "Dataset", "Configuration", "Run 1 Acc", "Run 2 Acc", "Std Dev", "Acc Range"]
            for i, h in enumerate(headers):
                cell = t6.cell(0, i)
                cell.text = h
                set_cell_background(cell, "1B365D")
                set_cell_margins(cell)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9.5)

            t6_data = [
                ["Gemma 2 2B", "Rotten Tomatoes", "strat_mixed_random_5_original", "0.5000", "0.1000", "0.2828", "0.4000"],
                ["DeepSeek R1 1.5B", "DBPedia 14", "strat_mixed_random_1_original", "0.5714", "0.2500", "0.2273", "0.3214"],
                ["DeepSeek R1 1.5B", "DBPedia 14", "strat_mixed_balanced_3_original", "0.2000", "0.2222", "0.0157", "0.0222"],
                ["Qwen 2.5 3B", "Rotten Tomatoes", "strat_mixed_random_3_similarity", "0.5000", "0.5000", "0.0000", "0.0000"]
            ]
            for r_idx, rdata in enumerate(t6_data, start=1):
                bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx, val in enumerate(rdata):
                    cell = t6.cell(r_idx, c_idx)
                    cell.text = val
                    set_cell_background(cell, bg)
                    set_cell_margins(cell)
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_table_borders(t6)

            # Insert Figure 4.8
            fig_p = doc.add_paragraph()
            fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p.paragraph_format.space_before = Pt(8)
            fig_p.paragraph_format.space_after = Pt(2)
            run = fig_p.add_run()
            run.add_picture("Run3/figures/16_reliability_variation.png", width=Inches(5.5))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run("Figure 4.8. Repeated-Run Reliability & Variance Analysis Across Valid Configurations")
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True

        elif header.startswith("4.9"):
            for ptext in body_paras:
                doc.add_paragraph(ptext)

    # Academic References Section
    h_ref = doc.add_heading(level=2)
    h_ref.paragraph_format.space_before = Pt(18)
    h_ref.paragraph_format.space_after = Pt(6)
    h_ref_run = h_ref.add_run("References")
    h_ref_run.font.name = 'Calibri'
    h_ref_run.font.size = Pt(15)
    h_ref_run.font.bold = True
    h_ref_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    refs = [
        "Demšar, J. (2006) ‘Statistical comparisons of classifiers over multiple data sets’, Journal of Machine Learning Research, 7, pp. 1–30.",
        "Dodge, J., Gururangan, S., Card, D., Schwartz, R. and Smith, N.A. (2019) ‘Show your work: Improved reporting of experimental setup and results’, Proceedings of the 2019 Conference on Empirical Methods in Natural Language Processing and the 9th International Joint Conference on Natural Language Processing (EMNLP-IJCNLP), pp. 2185–2194. doi: 10.18653/v1/D19-1224.",
        "Dror, R., Baumer, G., Shlomov, S. and Reichart, R. (2018) ‘The Hitchhiker’s guide to testing statistical significance in natural language processing’, Proceedings of the 56th Annual Meeting of the Association for Computational Linguistics (Volume 1: Long Papers), pp. 1383–1392. doi: 10.18653/v1/P18-1128.",
        "Reimers, N. and Gurevych, I. (2017) ‘Reporting score distributions makes a difference: Performance study of LSTM-networks for sequence tagging’, Proceedings of the 2017 Conference on Empirical Methods in Natural Language Processing, pp. 338–348. doi: 10.18653/v1/D17-1035."
    ]

    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_after = Pt(6)
        p_ref.paragraph_format.left_indent = Inches(0.5)
        p_ref.paragraph_format.first_line_indent = Inches(-0.5)
        r_ref = p_ref.add_run(ref)
        r_ref.font.size = Pt(10)

    output_filename = "dessertaion/Chapter_4_Results.docx"
    doc.save(output_filename)
    print(f"SUCCESS: Created Word document at '{output_filename}'!")

if __name__ == "__main__":
    build_docx()
